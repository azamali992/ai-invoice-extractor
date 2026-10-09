"""Turn results into tables and an Excel file a bookkeeper can open."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import Result

HEADER_FIELDS = [
    "vendor_name", "invoice_number", "invoice_date", "due_date", "currency",
    "subtotal", "tax", "shipping", "discount", "total",
    "customer_name", "vendor_tax_id", "vendor_address", "document_type",
]
REVIEW_FILL = PatternFill("solid", fgColor="FFF4CC")
ERROR_FILL = PatternFill("solid", fgColor="FADBD8")


def invoices_frame(results: list[Result]) -> pd.DataFrame:
    rows = []
    for r in results:
        row = {"file": r.filename, "status": r.status}
        if r.invoice:
            row.update({f: getattr(r.invoice, f) for f in HEADER_FIELDS})
            row["line_items"] = len(r.invoice.line_items)
        row["issues"] = "; ".join(i.message for i in r.issues) if r.issues else (r.error or "")
        rows.append(row)
    return pd.DataFrame(rows)


def line_items_frame(results: list[Result]) -> pd.DataFrame:
    rows = []
    for r in results:
        if not r.invoice:
            continue
        for n, item in enumerate(r.invoice.line_items, start=1):
            rows.append({
                "file": r.filename,
                "invoice_number": r.invoice.invoice_number,
                "line": n,
                **item.model_dump(),
            })
    return pd.DataFrame(rows, columns=["file", "invoice_number", "line", "description", "quantity", "unit_price", "amount"])


def to_excel(results: list[Result]) -> bytes:
    invoices = invoices_frame(results)
    lines = line_items_frame(results)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        invoices.to_excel(writer, sheet_name="Invoices", index=False)
        lines.to_excel(writer, sheet_name="Line items", index=False)

        ws = writer.sheets["Invoices"]
        cols = {name: idx + 1 for idx, name in enumerate(invoices.columns)}
        for row_idx, result in enumerate(results, start=2):
            if result.status == "ok":
                continue
            ws.cell(row=row_idx, column=cols["status"]).fill = ERROR_FILL if result.error else REVIEW_FILL
            for issue in result.issues:
                col = cols.get(issue.field)
                if col:
                    ws.cell(row=row_idx, column=col).fill = ERROR_FILL if issue.severity == "error" else REVIEW_FILL

        for sheet in writer.sheets.values():
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
            for col_cells in sheet.columns:
                width = max(len(str(c.value or "")) for c in col_cells)
                sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(width + 2, 10), 60)
    return buf.getvalue()


def to_csv(results: list[Result]) -> bytes:
    return invoices_frame(results).to_csv(index=False).encode()
