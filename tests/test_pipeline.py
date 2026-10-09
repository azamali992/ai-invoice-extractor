"""Tests run without any API key: a stub provider returns known values."""

import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from cli import score  # noqa: E402
from extractor.checks import check_invoice, pdf_text  # noqa: E402
from extractor.export import to_excel  # noqa: E402
from extractor.pipeline import process  # noqa: E402
from extractor.schema import Invoice, LineItem  # noqa: E402

SAMPLES = ROOT / "samples"
EXPECTED = json.loads((SAMPLES / "expected.json").read_text())


def truth_invoice(name: str, **overrides) -> Invoice:
    t = dict(EXPECTED[name])
    t.pop("should_flag")
    t.setdefault("vendor_address", None)
    t.update(document_type="invoice", vendor_address=t["vendor_address"], **overrides)
    return Invoice(**t)


class StubProvider:
    name = "stub"

    def __init__(self, **overrides):
        self.overrides = overrides

    def extract(self, data, mime_type):
        name = self.current
        return truth_invoice(name, **self.overrides.get(name, {}))


def run(name, provider=None):
    provider = provider or StubProvider()
    provider.current = name
    return process(name, (SAMPLES / name).read_bytes(), provider)


@pytest.mark.parametrize("name", [n for n, t in EXPECTED.items() if not t["should_flag"]])
def test_clean_documents_pass(name):
    r = run(name)
    assert r.status == "ok", [i.message for i in r.issues]


def test_wrong_subtotal_is_flagged():
    r = run("03_wrong_subtotal.pdf")
    assert r.status == "needs review"
    assert any(i.field == "subtotal" and i.severity == "error" for i in r.issues)


def test_text_pdf_is_grounded_scan_is_not():
    assert run("01_us_saas_invoice.pdf").grounded
    assert not run("05_scanned_invoice.pdf").grounded
    assert pdf_text((SAMPLES / "05_scanned_invoice.pdf").read_bytes()) == ""


def test_invented_total_is_caught_by_grounding_and_math():
    stub = StubProvider(**{"01_us_saas_invoice.pdf": {"total": 2500.00}})
    r = run("01_us_saas_invoice.pdf", stub)
    messages = " ".join(i.message for i in r.issues)
    assert "does not appear on the document" in messages
    assert any(i.field == "total" and i.severity == "error" for i in r.issues)


def test_line_math_and_dates():
    inv = Invoice(document_type="invoice", vendor_name="X", vendor_address=None, vendor_tax_id=None,
                  customer_name=None, invoice_number="1", invoice_date="2026-09-10", due_date="2026-09-01",
                  currency="USD", subtotal=30.0, tax=0.0, shipping=None, discount=None, total=30.0,
                  line_items=[LineItem(description="a", quantity=3, unit_price=10.0, amount=20.0)])
    fields = {i.field for i in check_invoice(inv)}
    assert {"line_items[1]", "subtotal", "due_date"} <= fields


def test_missing_required_fields_warn():
    inv = truth_invoice("01_us_saas_invoice.pdf", invoice_number=None, vendor_name=None)
    warned = {i.field for i in check_invoice(inv) if i.severity == "warning"}
    assert {"invoice_number", "vendor_name"} <= warned


def test_unsupported_file_type():
    r = process("notes.docx", b"x", StubProvider())
    assert r.status == "failed" and "unsupported" in r.error


def test_provider_error_does_not_crash_batch():
    class Broken:
        name = "broken"

        def extract(self, data, mime):
            raise TimeoutError("model timed out")

    r = process("01_us_saas_invoice.pdf", (SAMPLES / "01_us_saas_invoice.pdf").read_bytes(), Broken())
    assert r.status == "failed" and "timed out" in r.error


def test_excel_export_has_both_sheets_and_highlights_flags():
    results = [run(n) for n in sorted(EXPECTED)]
    wb = load_workbook(io.BytesIO(to_excel(results)))
    assert wb.sheetnames == ["Invoices", "Line items"]
    ws = wb["Invoices"]
    assert ws.max_row == len(results) + 1
    statuses = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2) for r in range(2, ws.max_row + 1)}
    assert statuses["03_wrong_subtotal.pdf"].fill.fgColor.rgb.endswith("FFF4CC")
    assert wb["Line items"].max_row - 1 == sum(len(t["line_items"]) for t in EXPECTED.values())


def test_score_against_ground_truth():
    results = [run(n) for n in sorted(EXPECTED)]
    report = score(results, EXPECTED)
    assert report["field_accuracy"] == 1.0
    assert report["flag_correct"] == f"{len(EXPECTED)}/{len(EXPECTED)}"
